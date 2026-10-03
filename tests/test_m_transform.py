import datetime
import json
from decimal import Decimal
from pathlib import Path

import duckdb
import openpyxl
import pytest

from sp_profiler import dataflow as df
from sp_profiler import mapping_load as ml
from sp_profiler import transforms as tf
from sp_profiler.cli import main
from sp_profiler.m2sql import Translator, register_udfs, stage_table, stage_value
from sp_profiler.mparse import MParseError, parse_expr
from sp_profiler.store import open_db

POLICIES_M = '''let
  Source = Excel.Workbook(Web.Contents("https://contoso.sharepoint.com/x.xlsx"), null, true),
  Sheet = Source{[Item="Data",Kind="Sheet"]}[Data],
  #"Promoted Headers" = Table.PromoteHeaders(Sheet, [PromoteAllScalars=true]),
  Renamed = Table.RenameColumns(#"Promoted Headers",{{"Policy No", "PolicyNumber"}, {"Insured Name", "Insured"}}),
  Filled = Table.FillDown(Renamed,{"Region"}),
  Typed = Table.TransformColumnTypes(Filled,{{"Gross Premium", type number}, {"Inception Date", type date}, {"Units", Int64.Type}}),
  Filtered = Table.SelectRows(Typed, each [PolicyNumber] <> null and [PolicyNumber] <> ""),
  Cleaned = Table.TransformColumns(Filtered,{{"Insured", each Text.Upper(Text.Trim(_)), type text}}),
  Added = Table.AddColumn(Cleaned, "Year", each Date.Year([Inception Date]), Int64.Type),
  Added2 = Table.AddColumn(Added, "Size", each if [Gross Premium] > 1000 then "Large" else "Small", type text),
  Merged = Table.NestedJoin(Added2, {"Currency"}, Currencies, {"Code"}, "Cur", JoinKind.LeftOuter),
  Expanded = Table.ExpandTableColumn(Merged, "Cur", {"Rate"}, {"FxRate"}),
  Gbp = Table.AddColumn(Expanded, "GrossGBP", each [Gross Premium] * [FxRate], type number),
  Cols = Table.SelectColumns(Gbp,{"PolicyNumber","Insured","Year","Size","GrossGBP","Region"})
in Cols'''
CURRENCIES_M = '#table({"Code","Rate"}, {{"GBP", 1.0}, {"EUR", 0.85}, {"USD", 0.79}})'
DOC = "section Section1;\r\nshared Currencies = " + CURRENCIES_M + ";\r\nshared Policies = " + POLICIES_M + ";\r\n"
MODEL = {"name": "Bordereau", "pbi:mashup": {"document": DOC}, "entities": [
    {"name": "Currencies", "attributes": [{"name": "Code", "dataType": "string"}, {"name": "Rate", "dataType": "double"}]},
    {"name": "Policies", "attributes": [
        {"name": "PolicyNumber", "dataType": "string"}, {"name": "Insured", "dataType": "string"},
        {"name": "Year", "dataType": "int64"}, {"name": "Size", "dataType": "string"},
        {"name": "GrossGBP", "dataType": "decimal"}, {"name": "Region", "dataType": "string"}]}]}


def test_parser_basics():
    ast = parse_expr('let #"A b" = Table.X(S, {{"a","b"}}), C = each [Col Name] > 1 and not [Z] in C')
    assert ast["steps"][0][0] == "A b"
    e = parse_expr('each if [Gross Premium] > 1 then "x" else "y"')
    assert e["t"] == "each" and e["body"]["t"] == "if"
    assert parse_expr('Source{[Item="Data",Kind="Sheet"]}[Data]')["t"] == "field"
    with pytest.raises(MParseError):
        parse_expr("let a = ")


def run_pipeline(M, entity, raw):
    p = Translator(M).translate(entity)
    cur = duckdb.connect().cursor()
    register_udfs(cur)
    n = len(next(iter(raw.values())))
    stage_table(cur, p.inputs, [(i + 2, *[stage_value(raw[c][i]) for c in p.inputs]) for i in range(n)])
    return p, cur.execute(p.sql()).fetchall(), [d[0] for d in cur.description]


def test_translate_full_policies_pipeline():
    raw = {"Policy No": ["P1", "P2", None, "P4"], "Insured Name": [" acme ", "beta", "x", "delta"],
           "Region": ["North", None, "South", None], "Gross Premium": [1200.5, "1,000", 1, 50],
           "Inception Date": [datetime.datetime(2024, 1, 5), "06/02/2024", None, "2023-12-31"],
           "Units": [3, "x", 1, 2.0], "Currency": ["GBP", "EUR", "USD", "GBP"]}
    p, rows, cols = run_pipeline({"Currencies": CURRENCIES_M, "Policies": POLICIES_M}, "Policies", raw)
    assert p.complete and set(p.inputs) == set(raw) and p.references["Currencies"]["kind"] == "query"
    got = {r[cols.index("PolicyNumber")]: dict(zip(cols, r)) for r in rows}
    assert set(got) == {"P1", "P2", "P4"}                       # blank policy filtered out
    assert got["P1"]["Insured"] == "ACME" and got["P1"]["Size"] == "Large" and got["P1"]["Year"] == 2024
    assert got["P2"]["GrossGBP"] == pytest.approx(850.0) and got["P2"]["Size"] == "Small"   # "1,000" parsed; EUR 0.85
    assert got["P2"]["Region"] == "North" and got["P4"]["Region"] == "South"                # fill down
    assert got["P1"]["__row"] == 2


def test_unsupported_step_is_reported_not_skipped():
    m = {"E": 'let S = Excel.Workbook(File.Contents("x"), null, true), P = Table.PromoteHeaders(S), '
              'G = Table.Group(P, {"a"}, {{"n", each List.Count([b]), type number}}), '
              'R = Table.RenameColumns(G, {{"a","A"}}) in R'}
    p = Translator(m).translate("E")
    assert not p.complete
    bad = [s for s in p.steps if not s["ok"]]
    assert [s["name"] for s in bad] == ["G", "R"] and "Table.Group" in bad[0]["error"]
    assert "depends" in bad[1]["error"] or "G" in bad[1]["error"] or "not" in bad[1]["error"]


def test_lookup_needing_data_and_binding():
    m = {"Fx": 'let S = Excel.Workbook(File.Contents("fx.xlsx"), null, true), D = S{[Item="Fx"]}[Data], '
               'P = Table.PromoteHeaders(D) in P',
         "E": 'let S = Excel.Workbook(File.Contents("x"), null, true), P = Table.PromoteHeaders(S), '
              'M = Table.NestedJoin(P, {"ccy"}, Fx, {"code"}, "F", JoinKind.LeftOuter), '
              'X = Table.ExpandTableColumn(M, "F", {"rate"}) in X'}
    p = Translator(m).translate("E")
    assert not p.complete and p.references["Fx"]["kind"] == "missing"
    con = duckdb.connect()
    con.execute("CREATE TABLE ref_fx (code VARCHAR, rate VARCHAR)")
    con.execute("INSERT INTO ref_fx VALUES ('GBP', '1'), ('EUR', '0.85')")
    tr = Translator(m, bindings={"Fx": "ref_fx"})
    p = tr.translate("E")
    assert p.complete and p.references["Fx"] == {"kind": "binding", "table": "ref_fx"}
    cur = con.cursor()
    register_udfs(cur)
    stage_table(cur, p.inputs, [(2, "GBP"), (3, "EUR"), (4, "JPY")])
    out = sorted(cur.execute(p.sql()).fetchall())
    assert out == [(2, "GBP", "1"), (3, "EUR", "0.85"), (4, "JPY", None)]     # (__row, ccy, rate); left join keeps JPY


def make_book(path: Path, headers, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Bordereau for March"])
    ws.append(headers)
    for r in rows:
        ws.append(r)
    wb.save(path)


@pytest.fixture
def env(tmp_path):
    src = tmp_path / "src"
    src.mkdir()
    make_book(src / "a.xlsx", ["Policy No", "Insured Name", "Region", "Gross Premium", "Inception Date", "Units", "Currency"],
              [["P1", " acme ", "North", 1200.5, datetime.datetime(2024, 1, 5), 3, "GBP"],
               ["P2", "beta", None, "1,000", "06/02/2024", 2, "EUR"],
               [None, None, None, None, None, None, None]])
    make_book(src / "b.xlsx", ["PolicyNumber", "Insured", "Region", "Premium (Gross)", "InceptionDate", "Units", "Ccy"],
              [["P3", "gamma", "South", 50, datetime.datetime(2023, 12, 1), 1, "USD"]])
    db = tmp_path / "t.duckdb"
    assert main(["local", "--path", str(src), "--db", str(db)]) == 0
    con = open_db(str(db))
    did, _ = df.store_dataflow(con, df.parse_model_json(json.dumps(MODEL)))
    return con, did


def map_inputs(con, did, entity="Policies"):
    plan = tf.resolve(con, did, entity)
    assert plan["mode"] == "m", plan["status"]
    for lh, nh in con.execute("select layout_hash, norm_headers from layouts").fetchall():
        sugg = df.suggest_mapping(nh, plan["inputs"])
        ml.save_mapping(con, entity, lh, [{"norm_header": h, "attribute": s["attribute"]} for h, s in sugg.items()],
                        plan["inputs"], "input")
    return plan


def test_m_mode_load_end_to_end(env):
    con, did = env
    plan = tf.resolve(con, did, "Policies")
    assert plan["use_m"] and plan["status"] == "complete" and "Currency" in plan["inputs"]
    map_inputs(con, did)
    # Currencies is a lookup *and* an entity; the lookup resolves from its inline #table
    r = ml.load_entity(con, did, "Policies")
    assert r["mode"] == "m" and r["sheets_loaded"] == 2 and r["rows"] == 3, r
    rows = {x[0]: x for x in con.execute(
        "select PolicyNumber, Insured, Year, Size, GrossGBP, Region, _excel_row, _coerce_errors from df_policies").fetchall()}
    assert rows["P1"][1:4] == ("ACME", 2024, "Large") and rows["P1"][6] == 3
    assert rows["P2"][4] == Decimal("850") and rows["P2"][5] == "North"
    assert rows["P3"][4] == Decimal("39.5") and rows["P3"][3] == "Small"       # layout B, USD 0.79
    assert ml.load_entity(con, did, "Policies")["sheets_planned"] == 0          # idempotent


def test_blocked_until_accept_partial_or_override(env):
    con, did = env
    bad_m = POLICIES_M.replace("Cols = Table.SelectColumns(Gbp,{\"PolicyNumber\",\"Insured\",\"Year\",\"Size\",\"GrossGBP\",\"Region\"})",
                               "Cols = Table.Pivot(Gbp, {\"a\"}, \"b\", \"c\")")
    model = json.loads(json.dumps(MODEL))
    model["name"] = "Bad"
    model["pbi:mashup"]["document"] = "section Section1;\r\nshared Currencies = " + CURRENCIES_M + ";\r\nshared Policies = " + bad_m + ";\r\n"
    did2, _ = df.store_dataflow(con, df.parse_model_json(json.dumps(model)))
    plan = tf.resolve(con, did2, "Policies")
    assert not plan["pipeline"].complete and plan["mode"] == "attribute" and plan["status"] == "disabled"
    tf.save_settings(con, did2, "Policies", True, "", False, [])
    assert tf.resolve(con, did2, "Policies")["status"].startswith("blocked")
    with pytest.raises(ValueError):
        ml.load_entity(con, did2, "Policies")
    # hand-written SQL override
    tf.save_settings(con, did2, "Policies", True,
                     "SELECT \"Policy No\" AS PolicyNumber, upper(trim(\"Insured Name\")) AS Insured, __row FROM _stg", False,
                     ["Policy No", "Insured Name"])
    plan = tf.resolve(con, did2, "Policies")
    assert plan["status"] == "override" and plan["mode"] == "m"
    for lh, nh in con.execute("select layout_hash, norm_headers from layouts").fetchall():
        sugg = df.suggest_mapping(nh, plan["inputs"])
        ml.save_mapping(con, "Policies", lh, [{"norm_header": h, "attribute": s["attribute"]} for h, s in sugg.items()],
                        plan["inputs"], "input")
    r = ml.load_entity(con, did2, "Policies")
    assert r["sheets_loaded"] >= 1 and r["rows"] >= 2
    assert con.execute("select count(*) from df_policies where Insured = 'ACME'").fetchone()[0] == 1


def test_reference_import_and_binding(env, tmp_path):
    con, did = env
    csv = tmp_path / "fx.csv"
    csv.write_text("Code,Rate\nGBP,1\nEUR,0.85\n")
    info = tf.import_reference(con, str(csv))
    assert info["table"] == "ref_fx" and info["rows"] == 2
    xl = tmp_path / "cur.xlsx"
    make_book(xl, ["Code", "Rate"], [["GBP", 1], ["EUR", 0.85]])
    assert tf.import_reference(con, str(xl), "ref_cur")["rows"] == 2
    with pytest.raises(ValueError):
        tf.import_reference(con, str(csv), "bad name;drop")


def test_pipeline_runs_in_worker_threads(env):
    """Regression: Python UDFs segfaulted DuckDB when executed from non-main threads."""
    import threading
    con, did = env
    map_inputs(con, did)
    results = {}

    def work():
        cur = con.cursor()
        results["r"] = ml.load_entity(cur, did, "Policies")
        results["n"] = tf.stage_sample(cur, tf.resolve(cur, did, "Policies"), "Policies",
                                       con.execute("select layout_hash from layouts limit 1").fetchone()[0])
    t = threading.Thread(target=work)
    t.start()
    t.join()
    assert results["r"]["rows"] == 3 and results["n"] >= 1
