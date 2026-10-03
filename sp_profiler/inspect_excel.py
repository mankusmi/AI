"""Open an Excel file and extract per-sheet structure: header row, headers, dimensions."""
from __future__ import annotations

import datetime
import hashlib
import re
import zipfile
from dataclasses import dataclass, field
from typing import Optional

EXCEL_EXTS = {".xlsx", ".xlsm", ".xltx", ".xltm"}
LEGACY_EXTS = {".xls", ".xlsb"}          # recognised, but not opened (no openpyxl support)


@dataclass
class SheetInfo:
    name: str
    state: str                       # visible / hidden / veryHidden
    max_row: int
    max_col: int
    header_row: Optional[int]        # 1-based, None if no header detected
    headers: list[str]               # raw header text
    norm_headers: list[str]          # normalised, order preserved
    data_rows: int                   # approx: max_row - header_row
    layout_hash: str = ""            # ordered normalised headers
    set_hash: str = ""               # same headers, order ignored

    @property
    def is_data_sheet(self) -> bool:
        return self.header_row is not None


@dataclass
class WorkbookInfo:
    status: str = "ok"               # ok / unsupported / encrypted / corrupt / error
    error: str = ""
    sheets: list[SheetInfo] = field(default_factory=list)
    sha256: str = ""
    defined_names: int = 0
    has_macros: bool = False
    warnings: list[str] = field(default_factory=list)


def normalise_header(text: object) -> str:
    s = re.sub(r"[^0-9a-z]+", " ", str(text).casefold())
    return re.sub(r"\s+", " ", s).strip()


def _short_hash(parts: list[str]) -> str:
    return hashlib.sha1("\x1f".join(parts).encode()).hexdigest()[:8]


def sha256_file(path, chunk: int = 1 << 20) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as fh:
        while block := fh.read(chunk):
            h.update(block)
    return h.hexdigest()


def _header_like(c) -> tuple[bool, bool]:
    """(counts as a header cell, is text). Years (2024) and dates count too: month/year columns are common."""
    if isinstance(c, str):
        return bool(c.strip()), bool(c.strip())
    if isinstance(c, bool):
        return False, False
    if isinstance(c, (int, float)):
        return float(c).is_integer() and 1900 <= float(c) <= 2100, False
    if isinstance(c, (datetime.datetime, datetime.date)):
        return True, False
    return False, False


def cell_text(c) -> str:
    """Header text for a cell: midnight datetimes as dates, 2024.0 as '2024'."""
    if c is None:
        return ""
    if isinstance(c, datetime.datetime):
        return c.date().isoformat() if c.time() == datetime.time(0) else c.isoformat(sep=" ")
    if isinstance(c, datetime.date):
        return c.isoformat()
    if isinstance(c, float) and c.is_integer():
        return str(int(c))
    return str(c).strip()


def detect_header_row(rows: list[tuple], min_headers: int = 3) -> Optional[int]:
    """Index (0-based) of the most header-like row among the scanned rows.

    Bordereaux often carry title/logo/summary rows above the real header, so we pick the first row with the
    most header-like cells (text, plus years and dates), requiring at least ``min_headers`` of them, at least
    one of them text, and mostly-unique values.
    """
    best, best_score = None, 0
    for i, row in enumerate(rows):
        flags = [_header_like(c) for c in row]
        cells = [c for c, (ok, _) in zip(row, flags) if ok]
        if len(cells) < min_headers or not any(is_text for _, is_text in flags):
            continue
        uniq = len({normalise_header(cell_text(c)) for c in cells})
        score = uniq if uniq >= 0.8 * len(cells) else 0
        if score > best_score:
            best, best_score = i, score
    return best


def _trim(row: tuple) -> list:
    r = list(row)
    while r and (r[-1] is None or (isinstance(r[-1], str) and not r[-1].strip())):
        r.pop()
    return r


def inspect_workbook(path, scan_rows: int = 50, min_headers: int = 3, include_hidden: bool = False,
                     exact_rows: bool = False) -> WorkbookInfo:
    from pathlib import Path
    path = Path(path)
    ext = path.suffix.lower()
    info = WorkbookInfo()
    try:
        info.sha256 = sha256_file(path)
    except OSError as e:
        info.status, info.error = "error", str(e)
        return info
    if ext in LEGACY_EXTS:
        info.status, info.error = "unsupported", f"{ext} not opened (convert to .xlsx)"
        return info
    if ext not in EXCEL_EXTS:
        info.status, info.error = "unsupported", "not an Excel file"
        return info

    import openpyxl
    try:
        with zipfile.ZipFile(path) as z:
            names = set(z.namelist())
        info.has_macros = "xl/vbaProject.bin" in names
    except zipfile.BadZipFile:
        # Password-protected workbooks are OLE2 containers, not zips
        with open(path, "rb") as fh:
            magic = fh.read(8)
        if magic.startswith(b"\xd0\xcf\x11\xe0"):
            info.status, info.error = "encrypted", "password-protected or legacy OLE container"
        else:
            info.status, info.error = "corrupt", "not a valid xlsx zip"
        return info

    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as e:  # openpyxl raises many types for damaged files
        info.status, info.error = "corrupt", f"{type(e).__name__}: {e}"
        return info
    try:
        info.defined_names = len(wb.defined_names)
        for ws in wb.worksheets:
            try:
                info.sheets.append(_inspect_sheet(ws, scan_rows, min_headers, include_hidden, exact_rows))
            except Exception as e:
                info.status, info.error = "error", f"sheet {ws.title!r}: {type(e).__name__}: {e}"
    finally:
        wb.close()
    data_names = [s.name for s in info.sheets if s.is_data_sheet]
    if data_names and info.status == "ok":
        n = _uncached_formulas(path, data_names)
        if n:
            info.warnings.append(f"{n} formula cell(s) have no cached value (workbook saved without calculating): "
                                 "they will read as empty. Open and re-save the file in Excel.")
    return info


def _uncached_formulas(path, sheet_names: list[str], limit: int = 300) -> int:
    """Count formula cells in the first rows whose calculated value was never stored in the file."""
    import openpyxl
    try:
        wf = openpyxl.load_workbook(path, read_only=True, data_only=False)
        wv = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception:
        return 0
    n = 0
    try:
        for name in sheet_names:
            for rf, rv in zip(wf[name].iter_rows(max_row=limit, values_only=True),
                              wv[name].iter_rows(max_row=limit, values_only=True)):
                n += sum(1 for f, v in zip(rf, rv) if isinstance(f, str) and f.startswith("=") and v is None)
    except Exception:
        return n
    finally:
        wf.close()
        wv.close()
    return n


def _count_rows(ws, first_data_row: int) -> int:
    return sum(1 for r in ws.iter_rows(min_row=first_data_row, values_only=True) if any(v is not None for v in r))


def _inspect_sheet(ws, scan_rows: int, min_headers: int, include_hidden: bool = False,
                   exact_rows: bool = False) -> SheetInfo:
    state = getattr(ws, "sheet_state", "visible")
    if getattr(ws, "max_row", None) is None:
        ws.reset_dimensions()
    max_row, max_col = ws.max_row or 0, ws.max_column or 0
    if state != "visible" and not include_hidden:
        return SheetInfo(ws.title, state, max_row, max_col, None, [], [], 0)    # hidden: recorded, not profiled
    head = [tuple(r) for r in ws.iter_rows(min_row=1, max_row=scan_rows, values_only=True)]
    idx = detect_header_row(head, min_headers)
    if idx is None:
        return SheetInfo(ws.title, state, max_row, max_col, None, [], [], 0)
    raw = _trim(head[idx])
    raw_txt = [cell_text(c) for c in raw]
    norm = [normalise_header(c) for c in raw_txt]
    # Blank header cells between named ones still occupy a column position
    norm = [n if n else f"<blank{i + 1}>" for i, n in enumerate(norm)]
    return SheetInfo(
        name=ws.title, state=state, max_row=max_row, max_col=max_col,
        header_row=idx + 1, headers=raw_txt, norm_headers=norm,
        data_rows=_count_rows(ws, idx + 2) if exact_rows else max(0, max_row - (idx + 1)),
        layout_hash=_short_hash(norm), set_hash=_short_hash(sorted(norm)),
    )
