"""Per-entity Power Query transformation settings and the pipeline that loading actually runs."""
from __future__ import annotations

from typing import Optional

from .m2sql import Pipeline, Translator, register_udfs
from .store import read_blob
from .mapping_load import date_order_for_culture, table_for


def dataflow_queries(con, dataflow_id: str) -> dict[str, str]:
    return dict(con.execute("SELECT name, m_query FROM dataflow_queries WHERE dataflow_id = ?", [dataflow_id]).fetchall())


def bindings(con, dataflow_id: str, coverholder: str = "") -> dict[str, str]:
    """Lookup query -> table. Bindings for the named coverholder override the shared ('') ones."""
    out: dict[str, str] = {}
    rows = con.execute("SELECT query_name, table_name, coverholder FROM query_bindings WHERE dataflow_id = ? "
                       "AND coverholder IN ('', ?) ORDER BY coverholder", [dataflow_id, coverholder or ""]).fetchall()
    for name, table, _ in rows:                    # '' sorts first, so the coverholder's own binding wins
        out[name] = table
    return out


def coverholder_bindings(con, dataflow_id: str) -> dict[str, dict[str, str]]:
    out: dict[str, dict[str, str]] = {}
    for name, table, ch in con.execute("SELECT query_name, table_name, coverholder FROM query_bindings WHERE dataflow_id = ? "
                                       "AND coverholder <> ''", [dataflow_id]).fetchall():
        out.setdefault(ch, {})[name] = table
    return out


def entity_tables(con, dataflow_id: str) -> dict[str, str]:
    """Entities of this dataflow that already have a loaded ``df_`` table (usable as lookup sources)."""
    out = {}
    for (ent,) in con.execute("SELECT entity FROM dataflow_entities WHERE dataflow_id = ?", [dataflow_id]).fetchall():
        tbl = table_for(con, ent)
        if con.execute("SELECT 1 FROM information_schema.tables WHERE table_name = ?", [tbl]).fetchone():
            out[ent] = tbl
    return out


def settings(con, dataflow_id: str, entity: str) -> Optional[dict]:
    r = con.execute("SELECT use_m, override_sql, accept_partial, extra_inputs, date_order FROM entity_transforms "
                    "WHERE dataflow_id = ? AND entity = ?", [dataflow_id, entity]).fetchone()
    return None if not r else {"use_m": r[0], "override_sql": r[1] or "", "accept_partial": bool(r[2]),
                               "extra_inputs": r[3] or [], "date_order": r[4] or "auto"}


def save_settings(con, dataflow_id: str, entity: str, use_m: bool, override_sql: str, accept_partial: bool,
                  extra_inputs: list[str], date_order: str = "auto") -> None:
    if date_order not in ("auto", "DMY", "MDY"):
        raise ValueError("date_order must be auto, DMY or MDY")
    con.execute("DELETE FROM entity_transforms WHERE dataflow_id = ? AND entity = ?", [dataflow_id, entity])
    con.execute("INSERT INTO entity_transforms (entity, dataflow_id, use_m, override_sql, accept_partial, extra_inputs, "
                "date_order) VALUES (?, ?, ?, ?, ?, ?, ?)",
                [entity, dataflow_id, use_m, override_sql.strip(), accept_partial, extra_inputs, date_order])


def date_order(con, dataflow_id: str, entity: str) -> tuple[str, str, str]:
    """(order, culture, source): 'DMY'/'MDY' from the entity setting, else from the dataflow culture."""
    st = settings(con, dataflow_id, entity)
    row = con.execute("SELECT culture FROM dataflows WHERE dataflow_id = ?", [dataflow_id]).fetchone()
    culture = (row[0] if row else "") or ""
    if st and st["date_order"] in ("DMY", "MDY"):
        return st["date_order"], culture, "setting"
    return date_order_for_culture(culture), culture, "culture"


def translate(con, dataflow_id: str, entity: str, coverholder: str = "") -> Pipeline:
    return Translator(dataflow_queries(con, dataflow_id), bindings(con, dataflow_id, coverholder),
                      entity_tables(con, dataflow_id), date_order(con, dataflow_id, entity)[0]).translate(entity)


def resolve(con, dataflow_id: str, entity: str, coverholder: str = "") -> dict:
    """Decide what a load will run: {mode: 'm'|'attribute', sql, inputs, status, pipeline, ...}."""
    p = translate(con, dataflow_id, entity, coverholder)
    st = settings(con, dataflow_id, entity)
    use = (st["use_m"] if st else p.complete)            # no explicit choice: auto-on only if fully translated
    extra = (st or {}).get("extra_inputs", [])
    order, culture, order_src = date_order(con, dataflow_id, entity)
    out = {"pipeline": p, "settings": st, "translated": p.complete, "use_m": bool(use), "mode": "attribute",
           "date_order": order, "culture": culture, "date_order_source": order_src,
           "sql": None, "inputs": [], "status": "disabled", "generated_sql": p.sql() if len(p.ctes) > 1 else "",
           "dynamic_columns": p.dynamic_columns}
    if not use:
        return out
    inputs = list(dict.fromkeys(p.inputs + extra))
    if st and st["override_sql"]:
        out.update(mode="m", sql=st["override_sql"], inputs=inputs, status="override")
    elif p.complete:
        out.update(mode="m", sql=p.sql(), inputs=inputs, status="complete")
    elif st and st["accept_partial"] and len(p.ctes) > 1:
        out.update(mode="m", sql=p.sql(), inputs=inputs, status="partial (accepted)")
    else:
        out.update(status="blocked: " + (p.error or "some steps are not supported; edit the SQL or accept partial"))
    return out


def output_columns(cur, plan: dict) -> list[tuple[str, str]]:
    register_udfs(cur)
    from .m2sql import stage_table
    stage_table(cur, plan["inputs"], [])
    return [(r[0], r[1]) for r in cur.execute(f"SELECT column_name, column_type FROM (DESCRIBE {plan['sql']})").fetchall()]


def stage_sample(con, plan: dict, entity: str, layout_hash: str, sample_rows: int = 500) -> int:
    """Stage the first rows of the first stored sheet with this layout (input mapping applied). Returns row count."""
    from .m2sql import sheet_columns, stage_table, stage_value
    from .mapping_load import _norm_header_cells, _sheet_rows
    row = con.execute("""SELECT f.sha256, s.sheet, s.header_row FROM sheets s JOIN files f USING (run_id, rel_path)
                         JOIN file_blobs b ON b.sha256 = f.sha256
                         WHERE s.layout_hash = ? AND s.header_row IS NOT NULL LIMIT 1""", [layout_hash]).fetchone()
    if not row:
        raise LookupError("No stored file with that layout (was it profiled with 'keep file bytes'?)")
    mapping = dict(con.execute("SELECT norm_header, attribute FROM column_mappings WHERE entity = ? AND layout_hash = ? "
                               "AND kind = 'input'", [entity, layout_hash]).fetchall())
    header, rows = _sheet_rows(read_blob(con, row[0]), row[1], row[2])
    cols = sheet_columns(header, plan["inputs"], mapping, _norm_header_cells(header), plan.get("dynamic_columns", False))
    staged = []
    for k, r in rows:
        vals = [stage_value(r[i]) if i is not None and i < len(r) else None for _, i in cols]
        if all(v is None or not v.strip() for v in vals):
            continue
        staged.append((k, *vals))
        if len(staged) >= sample_rows:
            break
    stage_table(con, [n for n, _ in cols], staged)
    return len(staged)


def _ref_name(stem: str) -> str:
    import re
    return "ref_" + (re.sub(r"[^0-9a-zA-Z]+", "_", stem).strip("_").lower() or "table")


def import_reference(con, path: str, table_name: str = "", sheet: str = "", *, source_id: str = "", sha256: str = "") -> dict:
    """Load a CSV/Excel lookup file into ``ref_<name>`` (all columns VARCHAR) so a query can be bound to it.

    The import is remembered (``reference_tables``) so ``refresh_references`` can re-import it when the file changes.
    """
    import hashlib
    from pathlib import Path
    from .bulk import bulk_insert
    from .inspect_excel import detect_header_row
    from .m2sql import q, stage_value
    p = Path(path).expanduser()
    name = table_name or _ref_name(p.stem)
    if not name.replace("_", "").isalnum():
        raise ValueError("Table name may only contain letters, digits and underscores")
    if p.suffix.lower() == ".csv":
        con.execute(f"CREATE OR REPLACE TABLE {q(name)} AS SELECT * FROM read_csv(?, all_varchar = true, header = true)", [str(p)])
    elif p.suffix.lower() in (".xlsx", ".xlsm"):
        import openpyxl
        wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
        try:
            ws = wb[sheet] if sheet else wb.worksheets[0]
            all_rows = [tuple(r) for r in ws.iter_rows(values_only=True)]
        finally:
            wb.close()
        hi = detect_header_row(all_rows[:50], 2)
        if hi is None:
            raise ValueError("Could not find a header row")
        from .export import delta_safe_names           # unique, safe names even if the sheet repeats a header
        raw = [str(c).strip() if c is not None and str(c).strip() else f"Column{i + 1}" for i, c in enumerate(all_rows[hi])]
        hdr = delta_safe_names(raw) if len({h.lower() for h in raw}) != len(raw) else raw
        con.execute(f"CREATE OR REPLACE TABLE {q(name)} ({', '.join(q(h) + ' VARCHAR' for h in hdr)})")
        data = [[stage_value(v) for v in (r + (None,) * len(hdr))[:len(hdr)]] for r in all_rows[hi + 1:]
                if any(v is not None for v in r)]
        bulk_insert(con, name, hdr, data)
    else:
        raise ValueError("Only .csv, .xlsx and .xlsm reference files are supported")
    sha = sha256 or hashlib.sha256(p.read_bytes()).hexdigest()
    con.execute("DELETE FROM reference_tables WHERE table_name = ?", [name])
    con.execute("INSERT INTO reference_tables (table_name, source_id, source_path, sheet, sha256) VALUES (?, ?, ?, ?, ?)",
                [name, source_id or None, str(p) if not source_id else None, sheet, sha])
    return {"table": name, "rows": con.execute(f"SELECT count(*) FROM {q(name)}").fetchone()[0],
            "columns": [r[0] for r in con.execute(f"SELECT column_name FROM (DESCRIBE {q(name)})").fetchall()]}


def import_reference_stored(con, sha256: str, table_name: str = "", sheet: str = "") -> dict:
    """Import a lookup workbook that was profiled with its bytes stored (e.g. the mapping file in a coverholder folder)."""
    import tempfile
    from pathlib import Path
    row = con.execute("SELECT name, COALESCE(NULLIF(item_id, ''), location, rel_path) FROM v_files WHERE sha256 = ? LIMIT 1",
                      [sha256]).fetchone()
    if not row:
        raise LookupError("No profiled file with that content")
    with tempfile.TemporaryDirectory(prefix="sp_ref_") as tmp:
        path = Path(tmp) / row[0]
        path.write_bytes(read_blob(con, sha256))
        return import_reference(con, str(path), table_name or _ref_name(Path(row[0]).stem), sheet, source_id=row[1], sha256=sha256)


def refresh_references(con) -> list[dict]:
    """Re-import lookup tables whose source file changed since they were imported. Returns what was refreshed."""
    import hashlib
    from pathlib import Path
    refreshed = []
    for table, source_id, path, sheet, sha in con.execute(
            "SELECT table_name, source_id, source_path, sheet, sha256 FROM reference_tables").fetchall():
        try:
            if source_id:
                cur = con.execute("SELECT sha256 FROM v_files WHERE COALESCE(NULLIF(item_id, ''), location, rel_path) = ? "
                                  "AND sha256 <> '' LIMIT 1", [source_id]).fetchone()
                if cur and cur[0] != sha:
                    info = import_reference_stored(con, cur[0], table, sheet or "")
                    refreshed.append({"table": table, "rows": info["rows"], "reason": "source file changed"})
            elif path and Path(path).exists() and hashlib.sha256(Path(path).read_bytes()).hexdigest() != sha:
                info = import_reference(con, path, table, sheet or "")
                refreshed.append({"table": table, "rows": info["rows"], "reason": "source file changed"})
        except Exception as e:          # a broken lookup must not stop the pipeline; the stage using it will report
            refreshed.append({"table": table, "error": f"{type(e).__name__}: {e}"})
    return refreshed


def table_fingerprint(con, table: str) -> str:
    """Cheap content signature of a table (row count + order-independent row hash), for change detection."""
    from .m2sql import q
    n, h = con.execute(f"SELECT count(*), coalesce(bit_xor(hash(t)), 0) FROM {q(table)} t").fetchone()
    return f"{n}:{h}"
